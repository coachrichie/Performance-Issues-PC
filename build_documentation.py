from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "PC-Performance-Diagnose-Dokumentation.docx"

BLUE = "2F5DA8"
DARK = "1F2937"
LIGHT = "EAF1FB"
GRAY = "F3F4F6"
MUTED = "6B7280"
WARN = "FFF4DB"


def set_font(run, name="Arial", size=11, color="000000", bold=None, italic=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:ascii"), name)
    rpr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def cell_margins(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", 80), ("start", 120), ("bottom", 80), ("end", 120)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths[index] / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_margins(cell)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        shade(cell, LIGHT)
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        set_font(paragraph.add_run(header), size=10, color=BLUE, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            paragraph = cells[i].paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            set_font(paragraph.add_run(str(value)), size=9.5, color=DARK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.2
    set_font(p.add_run(text))


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.2
    set_font(p.add_run(text))


def add_code(doc, text):
    p = doc.add_paragraph(style="Code Block")
    set_font(p.add_run(text), name="Consolas", size=9, color="1F2937")


def add_note(doc, label, text, fill=GRAY):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    shade(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    set_font(p.add_run(label + " "), size=10.5, color=BLUE, bold=True)
    set_font(p.add_run(text), size=10.5, color=DARK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_paragraph(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    set_font(p.add_run(text), color=DARK)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(1)
section.header_distance = section.footer_distance = Inches(0.45)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.15

for name, size, color, before, after in [
    ("Heading 1", 16, BLUE, 18, 8),
    ("Heading 2", 13, BLUE, 12, 6),
    ("Heading 3", 11.5, DARK, 10, 4),
]:
    style = styles[name]
    style.font.name = "Arial"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

code_style = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
code_style.font.name = "Consolas"
code_style.font.size = Pt(9)
code_style.paragraph_format.left_indent = Inches(0.2)
code_style.paragraph_format.space_before = Pt(3)
code_style.paragraph_format.space_after = Pt(8)
code_style.paragraph_format.line_spacing = 1.0

header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.LEFT
set_font(header.add_run("PC Performance Diagnose | Support Toolkit Dokumentation"), size=9, color=MUTED)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
set_font(footer.add_run("Stand: 29. Juli 2026"), size=8.5, color=MUTED)


p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(64)
p.paragraph_format.space_after = Pt(6)
set_font(p.add_run("PC PERFORMANCE DIAGNOSE"), size=26, color=DARK, bold=True)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(10)
set_font(p.add_run("Vollständige Dokumentation, Bedienungsanleitung und Übergabehilfe"), size=15, color=BLUE)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(18)
set_font(
    p.add_run("Windows 10/11 x64 | IT-Support-Triage | HTML-, CSV-, ZIP- und Word-Dokumentation"),
    size=10.5,
    color=MUTED,
)

add_note(
    doc,
    "Zielbild:",
    "Der Support kopiert einen vorbereiteten Ordner auf den Kunden-PC, startet Performance Test, lässt freigegebene Tools installieren und erhält Abschlussbericht, Rohbewertung und kumulative Logdaten.",
    LIGHT,
)

add_table(
    doc,
    ["Dokument", "Version", "Zweck"],
    [["PC-Performance-Diagnose-Dokumentation.docx", "2026-07-29", "Bedienung, Technik, Übergabe und GitHub-Nutzung"]],
    [3300, 1600, 4460],
)

doc.add_page_break()

doc.add_heading("1. Überblick", 1)
add_paragraph(doc, "Das Toolkit unterstützt die strukturierte Untersuchung von Performance-Problemen auf Windows-PCs. Es bündelt Windows-Abfragen, Benchmark-Ergebnisse, Support-Tool-Logs und normalisierte Auswertungen in einem wiederholbaren Ablauf.")
add_bullet(doc, "Zielgruppe: IT-Support, Triage, Second-Level-Support, interne Technik")
add_bullet(doc, "Primärer Startpunkt für Kundenfälle: Performance Test.cmd")
add_bullet(doc, "Zentrale Ausgaben: Abschlussbericht.html, Rohbewertung.html, Records.csv und Run.zip")
add_bullet(doc, "Sicherheitsprinzip: Diagnose statt Reparaturautomatik")

doc.add_heading("2. Unterstützter Kundenablauf", 1)
add_number(doc, "Support kopiert den vorbereiteten Toolkit-Ordner auf den Ziel-PC.")
add_number(doc, "Support startet Performance Test.cmd.")
add_number(doc, "Windows fordert bei Bedarf UAC/Elevation an.")
add_number(doc, "Das Toolkit installiert freigegebene AutoInstall-Werkzeuge aus dem Unterordner.")
add_number(doc, "Das Toolkit sammelt Windows-, Storage-, Netzwerk-, Benchmark- und Support-Tool-Daten.")
add_number(doc, "Das Toolkit erzeugt Abschlussbericht, Rohbewertung, CSV und ZIP.")
add_number(doc, "Support prüft die Berichte und grenzt die wahrscheinlichen Ursachen ein.")
add_note(
    doc,
    "Wichtiger Hinweis:",
    "Das Toolkit speichert kein Administrator-Passwort. Windows steuert die eigentliche UAC-Abfrage. Eine technisch garantierte Sammel-Passworteingabe für alle möglichen Fremdinstaller ist unter Windows nicht sauber erzwingbar.",
    WARN,
)

doc.add_heading("3. Hauptstart und Technikpfad", 1)
add_table(
    doc,
    ["Datei", "Rolle"],
    [
        ["Performance Test.cmd", "Einfacher Kundenstart per Doppelklick."],
        ["Performance Test.ps1", "Fordert Elevation an und startet den Vollworkflow."],
        ["Start-Diagnose.ps1", "Orchestriert Installation, Sammler, Benchmarks, Support-Tools und Reporting."],
    ],
    [2700, 6660],
)
add_code(doc, '.\\Start-Diagnose.ps1 -InstallTools -RunStressTests -RunBenchmarks -RunSupportTools')

doc.add_heading("4. Ordnerstruktur", 1)
add_table(
    doc,
    ["Ordner/Datei", "Bedeutung"],
    [
        ["Config\\", "Tools, Schwellenwerte und Benchmark-Referenzen."],
        ["Scripts\\", "Install, Collection, Benchmark, Support-Tool und Reporting-Logik."],
        ["Tests\\", "Pester-Absicherung des Toolkits."],
        ["ToolkitPrograms\\AutoInstall\\", "Bevorzugter Ablageort für automatisch installierbare Werkzeuge."],
        ["ToolkitPrograms\\Optional\\", "Manuelle/optionale Werkzeuge ohne AutoInstall-Zwang."],
        ["ToolkitPrograms\\Installed\\", "Ablage für entpackte ZIP-Tools während eines Laufs."],
        ["BenchmarkImports\\", "Importordner für PCMark 10 XML und Unigine Heaven CSV."],
        ["Reports\\", "Ergebnisse vergangener Läufe."],
        ["Distributions\\", "Automatisch erstellte Kunden- und GitHub-Übergabeordner."],
    ],
    [2900, 6460],
)

doc.add_heading("5. Unterstützte Diagnosebereiche", 1)
add_table(
    doc,
    ["Bereich", "Beispiele"],
    [
        ["Windows", "OS, Build, Modell, RAM, Prozesse, Systemereignisse."],
        ["Storage", "Freier Speicher, Volumes, Disk-Health, I/O-Hinweise."],
        ["Network", "Adapter, IPv4, DNS, Gateway, Grundkontext ohne Mitschnitt."],
        ["Benchmarks", "WinSAT, PCMark 10, Unigine Heaven."],
        ["Support-Tools", "CrystalDiskInfo, Autoruns, Process Explorer."],
        ["Reporting", "Kumulative Normalisierung und Ursachenhinweise."],
    ],
    [2300, 7060],
)

doc.add_heading("6. Support-Tools und Logauswertung", 1)
add_table(
    doc,
    ["Tool", "Automatik im Projekt", "Auswertung"],
    [
        ["CrystalDiskInfo", "Kann aus ZIP bereitgestellt werden; Logexport via /CopyExit.", "Laufwerkszustand und Temperatur werden in den kumulativen Record-Stream übernommen."],
        ["Autoruns", "Kann aus ZIP bereitgestellt werden; CSV-Export per Autorunsc.", "Autostart-Gesamtzahl und unsigned/auffällige Einträge fließen in die Bewertung ein."],
        ["Process Explorer", "Wird für tiefe manuelle Prozesssicht gestartet.", "Launch wird protokolliert; keine offizielle CLI-Vollauswertung hinterlegt."],
    ],
    [1700, 2650, 5010],
)

doc.add_heading("7. Benchmark-Bewertung", 1)
add_paragraph(doc, "Das Toolkit kann Benchmark-Daten importieren, automatisierte Startversuche vorbereiten und Messergebnisse mit lokalen Referenzwerten vergleichen.")
add_bullet(doc, "WinSAT wird direkt über Windows abgefragt.")
add_bullet(doc, "PCMark 10 wird über XML-Exporte oder unterstützte Automation bewertet.")
add_bullet(doc, "Unigine Heaven wird über CSV-Exporte oder vorbereitete Automation bewertet.")
add_bullet(doc, "Die Vergleichslogik markiert Werte als InExpectedRange, UnderExpectedRange oder BelowMinimum.")

doc.add_heading("8. Berichte und Ausgabedateien", 1)
add_table(
    doc,
    ["Datei", "Zweck"],
    [
        ["Abschlussbericht.html", "Management-/Support-taugliche Kurzbewertung mit wahrscheinlichen Ursachen und nächsten Schritten."],
        ["Rohbewertung.html", "Technische Detailansicht aller normalisierten Datensätze."],
        ["Records.csv", "Kumulative tabellarische Weiterverarbeitung und Mehrlaufvergleich."],
        ["Run.zip", "Bündelung der Berichte, Rohdaten und Logs für Archiv oder Übergabe."],
    ],
    [2600, 6760],
)

doc.add_heading("9. Sicherheit und bewusste Grenzen", 1)
add_bullet(doc, "Keine automatische Bereinigung oder Deinstallation.")
add_bullet(doc, "Keine Registry-, Treiber- oder Energieplanänderungen.")
add_bullet(doc, "Keine Passwortspeicherung.")
add_bullet(doc, "Keine automatische Wireshark-Paketaufzeichnung.")
add_bullet(doc, "CHIP-Wrapper und ähnliche fragwürdige Installer bleiben ausgeschlossen.")
add_note(
    doc,
    "Datenschutz:",
    "Berichte und Rohdaten können Rechnernamen, Benutzernamen, Pfade, Prozessdaten, Ereignisdetails und Netzwerk-Kontext enthalten. Vor externer Weitergabe immer prüfen.",
    WARN,
)

doc.add_heading("10. Empfohlene zusätzliche Freeware", 1)
add_table(
    doc,
    ["Programm", "Nutzen im Support"],
    [
        ["CrystalDiskInfo", "Schnelle Health-/SMART-Sicht für Datenträgerprobleme."],
        ["Autoruns", "Erkennt überladene Autostarts und Treiber-/Shell-Hooks."],
        ["Process Explorer", "Tiefere Prozessanalyse als Standard-Task-Manager."],
    ],
    [2500, 6860],
)

doc.add_heading("11. GitHub- und Release-Nutzung", 1)
add_paragraph(doc, "Das Projekt kann über das Skript Create-Distributions.ps1 in zwei saubere Handoff-Ordner exportiert werden.")
add_table(
    doc,
    ["Ziel", "Ergebnis"],
    [
        ["Distributions\\Customer Toolkit\\", "Kunden-/Supportordner zum Kopieren auf einen Ziel-PC."],
        ["Distributions\\GitHub Repository\\", "Bereinigter Projektordner für Versionsverwaltung und Veröffentlichung."],
    ],
    [3400, 5960],
)
add_code(doc, ".\\Create-Distributions.ps1")

doc.add_heading("12. Bedienung auf dem Kunden-PC", 1)
add_number(doc, "Den Ordner Distributions\\Customer Toolkit oder den vorbereiteten Hauptordner lokal auf den Kunden-PC kopieren.")
add_number(doc, "Freigegebene Tools in ToolkitPrograms\\AutoInstall prüfen oder ergänzen.")
add_number(doc, "Performance Test.cmd starten.")
add_number(doc, "Windows-UAC bestätigen.")
add_number(doc, "Auf Abschluss des Laufes warten.")
add_number(doc, "Abschlussbericht.html öffnen.")
add_number(doc, "Für technische Details Rohbewertung.html und Records.csv prüfen.")
add_note(
    doc,
    "Praxisregel:",
    "Der Bericht liefert eine strukturierte Triage und Ursachenhinweise. Er ersetzt keine fachliche Bewertung durch den Support, reduziert aber die Suchzeit deutlich.",
    LIGHT,
)

doc.add_heading("13. Verifikation", 1)
add_code(doc, "Invoke-Pester -Script .\\Tests -PassThru")
add_code(doc, ".\\Start-Diagnose.ps1 -DryRun -InstallTools -RunStressTests -RunBenchmarks -RunSupportTools -OutputRoot .\\VerificationOutput")
add_code(doc, ".\\Create-Distributions.ps1")
add_bullet(doc, "Vor Übergabe immer erst Testsuite und Dry-Run durchführen.")
add_bullet(doc, "Word-Dokumentation nach inhaltlichen Änderungen neu erzeugen.")
add_bullet(doc, "Distributions-Ordner nach jeder strukturellen Änderung neu erzeugen.")

doc.add_heading("14. Kurz-Checkliste für den praktischen Einsatz", 1)
for item in [
    "ToolkitPrograms\\AutoInstall mit freigegebenen Dateien befüllen",
    "Optionale portable Hilfstools in ToolkitPrograms\\Optional ablegen",
    "Performance Test auf Zielsystem starten",
    "Abschlussbericht und Rohbewertung prüfen",
    "CSV und ZIP nur nach Datenschutzprüfung weitergeben",
    "Bei Benchmarkfragen Exporte in BenchmarkImports ablegen oder Automation nutzen",
]:
    add_bullet(doc, item)

doc.save(OUT)
print(OUT)
